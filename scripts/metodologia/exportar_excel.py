"""Export de texto del Excel maestro (workflow.md §14.7 y §14.9).

Uso: python -I scripts/metodologia/exportar_excel.py "docs/excel/Base de Datos Valpo Verde.xlsx" docs/excel/export
Requiere openpyxl. Genera un CSV UTF-8 por hoja del paquete (núcleo normativo, gobernanza,
menús e INDICES; PR-002 v3.0), con valores (no fórmulas), en un orden determinista, para
revisar los cambios del Excel como texto en git. No exporta los anexos operativos: no son
fuente metodológica y pueden contener datos personales (p. ej. USUARIOS). No modifica el Excel.
"""
import csv, os, re, sys
import openpyxl

src, out = sys.argv[1], sys.argv[2]
os.makedirs(out, exist_ok=True)
wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
HOJAS = ["DICCIONARIO_CAMPOS", "REGLAS_GENERALES", "REGLAS_INDICADORES", "MATRICES_CALCULO",
         "VERSION", "LISTAS", "INDICES"]
faltan = [h for h in HOJAS if h not in wb.sheetnames]
if faltan:
    sys.exit(f"Faltan hojas del paquete: {faltan}")
for ws in (wb[h] for h in HOJAS):
    rows = [["" if v is None else str(v).replace("\r\n", "\n") for v in r] for r in ws.iter_rows(values_only=True)]
    while rows and not any(rows[-1]):
        rows.pop()
    width = max((max((i + 1 for i, v in enumerate(r) if v), default=0) for r in rows), default=0)
    name = re.sub(r"[^A-Za-z0-9_]+", "_", ws.title).strip("_")
    with open(os.path.join(out, f"{name}.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        for r in rows:
            w.writerow(r[:width])
    print(f"{ws.title}: {len(rows)} filas -> {name}.csv")
